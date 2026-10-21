from __future__ import annotations
import os, re
from pathlib import Path
from typing import Literal
import gdown, requests, fitz
from docx import Document
from pptx import Presentation
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

load_dotenv()
BASE_DIR = Path(__file__).resolve().parent
MATERIALS_DIR = Path(os.getenv("MATERIALS_DIR", str(BASE_DIR / "materiais"))).resolve()
DRIVE_URL = os.getenv("GOOGLE_DRIVE_FOLDER_URL", "https://drive.google.com/drive/folders/1EUof9_ZX-FiLA1GTi8le5AQpccKH08je?usp=sharing")
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434").rstrip("/")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:3b")
MAX_FILE_MB = int(os.getenv("MAX_FILE_MB", "50"))
SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md"}
app = FastAPI(title="AprovaJá — IA local gratuita", version="1.1.0")
app.add_middleware(CORSMiddleware, allow_origins=["https://fpontualprog.github.io", "https://aprovajaconcurso.netlify.app", "http://localhost:8000", "http://127.0.0.1:8000", "http://localhost:5500", "http://127.0.0.1:5500"], allow_credentials=False, allow_methods=["GET", "POST"], allow_headers=["Content-Type"])

class GenerateRequest(BaseModel):
    topic: str = Field(min_length=3, max_length=300)
    kind: Literal["resumo", "mapa_mental", "questoes", "flashcards", "carrossel"] = "resumo"
    discipline: str = Field(default="Concursos públicos", max_length=120)
    quantity: int = Field(default=5, ge=1, le=20)
    exam_board: str = Field(default="CEBRASPE", max_length=80)

def extract_file(path: Path) -> list[dict]:
    chunks = []
    try:
        if path.suffix.lower() == ".pdf":
            with fitz.open(path) as pdf:
                for n, page in enumerate(pdf, 1):
                    t = page.get_text("text").strip()
                    if t: chunks.append({"text": t, "source": path.name, "page": n})
        elif path.suffix.lower() == ".docx":
            doc = Document(path)
            t = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            if t: chunks.append({"text": t, "source": path.name, "page": None})
            for i, table in enumerate(doc.tables, 1):
                t = "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in table.rows)
                if t.strip(): chunks.append({"text": t, "source": f"{path.name} (tabela {i})", "page": None})
        elif path.suffix.lower() == ".pptx":
            prs = Presentation(path)
            for n, slide in enumerate(prs.slides, 1):
                parts = [s.text for s in slide.shapes if getattr(s, "has_text_frame", False) and s.text.strip()]
                if parts: chunks.append({"text": "\n".join(parts), "source": path.name, "page": n})
        else:
            t = path.read_text(encoding="utf-8", errors="ignore").strip()
            if t: chunks.append({"text": t, "source": path.name, "page": None})
    except Exception as exc:
        print(f"[aviso] Falha ao ler {path.name}: {exc}")
    return chunks

def load_chunks() -> list[dict]:
    MATERIALS_DIR.mkdir(parents=True, exist_ok=True)
    result = []
    for p in MATERIALS_DIR.rglob("*"):
        if p.is_file() and p.suffix.lower() in SUPPORTED and p.stat().st_size <= MAX_FILE_MB * 1024 * 1024:
            result.extend(extract_file(p))
    return result

def split_into_passages(chunk: dict, max_chars: int = 1800) -> list[dict]:
    """Divide páginas/documentos longos em passagens menores para evitar diluir a relevância."""
    text = re.sub(r"\r", "", chunk.get("text", "")).strip()
    if len(text) <= max_chars:
        return [{**chunk, "text": text}] if text else []
    paragraphs = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    passages, current = [], ""
    for paragraph in paragraphs:
        # Parágrafo enorme: fragmenta em janelas com pequena sobreposição.
        pieces = [paragraph[i:i + max_chars] for i in range(0, len(paragraph), max_chars - 180)] if len(paragraph) > max_chars else [paragraph]
        for piece in pieces:
            if current and len(current) + len(piece) + 1 > max_chars:
                passages.append({**chunk, "text": current})
                current = piece
            else:
                current = (current + "\n" + piece).strip()
    if current:
        passages.append({**chunk, "text": current})
    return passages

def retrieve(topic: str, discipline: str, chunks: list[dict], limit: int = 6) -> list[dict]:
    if not chunks:
        return []
    passages = [p for c in chunks for p in split_into_passages(c)]
    if not passages:
        return []
    query = f"{discipline} {topic} conceitos definição legislação exemplos questões concurso"
    docs = [re.sub(r"\s+", " ", c["text"])[:4000] for c in passages]
    try:
        vec = TfidfVectorizer(max_features=60000, ngram_range=(1, 2), strip_accents="unicode", sublinear_tf=True)
        matrix = vec.fit_transform(docs + [query])
        scores = cosine_similarity(matrix[-1], matrix[:-1]).ravel()
        ranked = scores.argsort()[::-1]
        selected = []
        # Limite de duas passagens por arquivo para não lotar o contexto com uma única fonte.
        per_source = {}
        for idx in ranked:
            score = float(scores[idx])
            if score < 0.035:
                break
            source = passages[idx]["source"]
            if per_source.get(source, 0) >= 2:
                continue
            selected.append({**passages[idx], "relevance_score": round(score, 4)})
            per_source[source] = per_source.get(source, 0) + 1
            if len(selected) >= limit:
                break
        return selected
    except ValueError:
        return []

def make_prompt(req: GenerateRequest, sources: list[dict]) -> str:
    snippets = "\n\n".join(
        f"[FONTE {i}: {s['source']}" + (f", página/slide {s['page']}" if s.get("page") else "") + f", relevância {s.get('relevance_score', 'n/d')}]\n{s['text'][:2400]}"
        for i, s in enumerate(sources, 1)
    )
    if not snippets:
        snippets = "Nenhum trecho atingiu o nível mínimo de relevância nos materiais locais. Não alegue ter consultado apostilas. Produza somente o que puder sustentar com conhecimento geral e inclua aviso claro para conferir fontes oficiais."
    fmt = {
        "resumo": f"Faça um resumo de estudo sobre EXATAMENTE o tema solicitado. Inclua definição, regras/elementos centrais, distinções importantes, exemplo curto e pegadinhas de concurso. Não se desvie para outro assunto.",
        "mapa_mental": "Crie um mapa mental textual hierárquico, com ramos e sub-ramos reais e preenchidos. Organize em 4 a 7 ramos úteis, sem níveis vazios como 'Ramos' e 'Sub-ramos'. Use frases curtas, não repita o mesmo item e apenas inclua artigos de lei ou referências normativas que estejam nos trechos recuperados. Se não houver base para uma referência específica, omita-a em vez de adivinhar.",
        "questoes": f"Crie EXATAMENTE {req.quantity} questões inéditas sobre o tema e a disciplina solicitados, no estilo {req.exam_board}. Numere de 1 a {req.quantity}. Cada questão deve ter enunciado completo, alternativas quando apropriado, gabarito e comentário explicando o fundamento. Não repita questões e não troque o tema por assunto relacionado.",
        "flashcards": f"Crie EXATAMENTE {req.quantity} flashcards sobre o tema solicitado. Numere de 1 a {req.quantity}, cada um com Pergunta e Resposta curtas. Não acrescente cartões extras.",
        "carrossel": "Crie EXATAMENTE 7 slides numerados: 1 capa, 2 a 6 conteúdo didático, 7 fechamento/CTA. Cada slide deve conter título e texto curto, sobre o tema pedido."
    }[req.kind]
    return f"""Você é um professor rigoroso e revisor de material para concursos públicos brasileiros. Escreva em português brasileiro claro.

REQUISITOS OBRIGATÓRIOS
- Disciplina exata: {req.discipline}
- Tema exato: {req.topic}
- Formato: {req.kind}
- Banca de referência, se aplicável: {req.exam_board}
- Instrução específica: {fmt}
- Antes de responder, confirme mentalmente que cada seção trata diretamente do tema. Não substitua o tema por um caso isolado nem por um assunto apenas relacionado.
- Use primeiro os trechos recuperados, mas avalie sua pertinência. Um trecho de outro assunto NÃO deve servir de fundamento só porque veio da mesma pasta.
- REGRA CRÍTICA PARA DIREITO: não invente nem improvise números de artigos, leis, processos, decisões, súmulas ou jurisprudência. Só mencione referências normativas específicas que constem literalmente dos trechos recuperados. Se não houver trecho relevante, explique os conceitos sem fabricar citações ou precedentes. Nunca atribua à educação a Lei nº 8.080/1990 (ela é da saúde); nunca atribua moradia ao art. 194 (o direito à moradia está no art. 6º; art. 194 trata da seguridade social).
- Não invente fontes, citações, números de páginas, enunciados de lei ou jurisprudência. Só atribua uma afirmação a uma fonte quando ela estiver apoiada no trecho correspondente.
- Para Direito, diferencie texto legal, entendimento jurisprudencial e explicação didática. Se não tiver certeza de artigo, prazo ou entendimento atual, não chute: sinalize a necessidade de conferência em fonte oficial.
- Se os materiais não forem pertinentes ou forem insuficientes, diga isso com transparência e não finja que foram utilizados.
- Evite introduções genéricas, repetições e conclusões vagas. Priorize precisão, utilidade para prova e aderência ao pedido.
- Ao final, inclua “Pontos que exigem conferência” somente se houver incerteza real e “Fontes efetivamente utilizadas” apenas para trechos realmente aproveitados.
- Não produza listas com duplicações para preencher espaço. Para mapa mental, use hierarquia substantiva e direta, por exemplo: Direitos Sociais > Art. 6º; Direitos dos Trabalhadores > Arts. 7º a 11; Saúde > Art. 196; Educação > Art. 205. Não force este exemplo se o tema for outro.

MATERIAIS RECUPERADOS (podem ser irrelevantes; use com senso crítico):
{snippets}

Agora produza a resposta obedecendo todos os requisitos acima."""

def validate_content(req: GenerateRequest, content: str) -> list[str]:
    warnings = []
    if len(content.strip()) < 180:
        warnings.append("Resposta muito curta; revisar antes de usar.")
    if req.kind in {"questoes", "flashcards"}:
        # Detecta itens numerados no início de linhas; é um controle simples, não uma avaliação semântica.
        nums = set()
        for line in content.splitlines():
            match = re.match(r"^\s*(?:#{1,4}\s*)?(?:Questão\s*)?(\d{1,2})[.)\-:]\s+", line, re.I)
            if match:
                nums.add(int(match.group(1)))
        if len(nums) < req.quantity:
            warnings.append(f"Controle de quantidade: foram detectados {len(nums)} itens numerados; eram esperados {req.quantity}. Verifique e regenere se necessário.")
    if req.kind == "carrossel":
        slides = set(re.findall(r"(?im)^\s*(?:slide\s*)?(\d{1,2})\s*[:.)\-]", content))
        if len(slides) < 7:
            warnings.append(f"Controle de formato: foram detectados {len(slides)} slides numerados; eram esperados 7.")
    # Guardrails específicos para erros jurídicos conhecidos e citações suspeitas.
    legal_text = content.lower()
    if "direito constitucional" in req.discipline.lower() or any(k in req.topic.lower() for k in ("direito", "constituição", "constitucional", "lei", "artigo", "direitos sociais")):
        if re.search(r"lei\s*n[ºo.]?\s*8[.]080\s*/?\s*1990.{0,100}(educação|educacional|ldb)", legal_text, re.S):
            warnings.append("ALERTA JURÍDICO: possível atribuição incorreta da Lei nº 8.080/1990 à educação. Conferir fonte oficial antes de publicar.")
        if re.search(r"(moradia|direito à moradia).{0,100}art(?:igo|\.)?\s*194", legal_text, re.S):
            warnings.append("ALERTA JURÍDICO: possível associação incorreta entre moradia e art. 194 da CF. Conferir fonte oficial antes de publicar.")
        if re.search(r"(decisão|acórdão|processo|súmula|precedente).{0,80}\b(?:\d{1,2}[.]\d{3,}|\d{4}/\d{2,4})\b", legal_text, re.S):
            warnings.append("ALERTA JURÍDICO: há possível número de processo ou precedente que requer verificação oficial.")
    return warnings

@app.get("/health")
def health():
    try:
        r = requests.get(f"{OLLAMA_URL}/api/tags", timeout=3); r.raise_for_status()
        models = [m.get("name", "") for m in r.json().get("models", [])]
        reachable = True
        available = any(m == OLLAMA_MODEL or m.startswith(OLLAMA_MODEL + ":") for m in models)
        detail = {"models_installed": models, "configured_model": OLLAMA_MODEL}
    except Exception as exc:
        reachable, available, detail = False, False, {"error": str(exc), "configured_model": OLLAMA_MODEL}
    return {"status": "ok", "service": "AprovaJá IA local", "ollama_reachable": reachable, "model_available": available, "details": detail}

@app.get("/materials")
def materials():
    chunks = load_chunks()
    names = sorted({c["source"].split(" (tabela")[0] for c in chunks})
    return {"files_read": len(names), "chunks_indexed": len(chunks), "files": names, "directory": str(MATERIALS_DIR)}

@app.post("/sync-drive")
def sync_drive():
    MATERIALS_DIR.mkdir(parents=True, exist_ok=True)
    try:
        downloaded = gdown.download_folder(url=DRIVE_URL, output=str(MATERIALS_DIR), quiet=False, use_cookies=False, remaining_ok=True)
        if downloaded is None:
            raise HTTPException(status_code=502, detail="O Drive não retornou arquivos. Confira o link público e possíveis limites temporários de download.")
        chunks = load_chunks()
        names = sorted({c["source"].split(" (tabela")[0] for c in chunks})
        return {"downloaded_entries": len(downloaded), "files_read": len(names), "chunks_indexed": len(chunks), "directory": str(MATERIALS_DIR), "files": names, "note": "Cópias locais baixadas; os originais do Drive não foram alterados."}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Falha ao sincronizar o Drive: {exc}")

@app.post("/generate")
def generate(req: GenerateRequest):
    all_chunks = load_chunks()
    sources = retrieve(req.topic, req.discipline, all_chunks)
    prompt = make_prompt(req, sources)
    try:
        r = requests.post(
            f"{OLLAMA_URL}/api/chat",
            json={"model": OLLAMA_MODEL,
                  "messages": [
                      {"role": "system", "content": "Você produz materiais educacionais para concursos. Siga exatamente o tema, a disciplina e a quantidade solicitados. Não invente fontes nem complete lacunas jurídicas com palpites."},
                      {"role": "user", "content": prompt}],
                  "stream": False,
                  "options": {"temperature": 0.2, "top_p": 0.8}},
            timeout=600)
        if r.status_code == 404:
            raise HTTPException(status_code=503, detail=f"Modelo '{OLLAMA_MODEL}' não instalado. Execute: ollama pull {OLLAMA_MODEL}")
        r.raise_for_status()
        content = r.json().get("message", {}).get("content", "").strip()
        if not content:
            raise HTTPException(status_code=502, detail="O modelo local retornou conteúdo vazio.")
        refs = [{"source": s["source"], "page": s.get("page"), "relevance_score": s.get("relevance_score")} for s in sources]
        warnings = validate_content(req, content)
        if not refs:
            warnings.append("Nenhuma fonte local atingiu o limite de relevância. Revise todas as afirmações, especialmente legislação e gabaritos.")
        return {"content": content, "topic": req.topic, "discipline": req.discipline, "kind": req.kind,
                "model": OLLAMA_MODEL, "sources": refs, "source_count": len(refs),
                "warnings": warnings, "warning": "; ".join(warnings) if warnings else None,
                "retrieval": {"chunks_available": len(all_chunks), "passages_considered": sum(len(split_into_passages(c)) for c in all_chunks), "passages_selected": len(sources)}}
    except HTTPException:
        raise
    except requests.ConnectionError:
        raise HTTPException(status_code=503, detail="Ollama não está em execução. Instale e inicie o Ollama localmente.")
    except requests.Timeout:
        raise HTTPException(status_code=504, detail="O modelo local demorou demais. Tente um modelo menor.")
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Erro ao chamar o modelo local: {exc}")
